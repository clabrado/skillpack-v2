#!/usr/bin/env python3
"""ad-qa extractor: turn any target file into reviewable text (+ page PNGs for visual formats).

Usage: python3 -P extract.py <file-or-dir> <out_dir>
Writes <out_dir>/<name>.txt and, for visual formats, <out_dir>/<name>-pages/page-N.png.
Prints a one-line manifest per file. Never executes the input.
"""
import csv, json, os, shutil, subprocess, sys, zipfile
from pathlib import Path

PDFTOTEXT = shutil.which("pdftotext")
PDFTOPPM = shutil.which("pdftoppm")
SOFFICE = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
TEXT_EXT = {".txt", ".md", ".json", ".yaml", ".yml", ".toml", ".ini", ".log", ".xml", ".sql",
            ".py", ".ts", ".tsx", ".js", ".jsx", ".go", ".rs", ".java", ".sh", ".css", ".eml"}
IMG_EXT = {".png", ".jpg", ".jpeg", ".gif", ".webp"}


def run(cmd):
    return subprocess.run(cmd, capture_output=True, text=True, timeout=180)


def render_pdf(pdf: Path, out: Path, label: str = ""):
    if not PDFTOPPM:
        return "no pdftoppm"
    d = out / f"{label or pdf.name}-pages"
    d.mkdir(parents=True, exist_ok=True)
    run([PDFTOPPM, "-png", "-r", "70", str(pdf), str(d / "page")])
    return f"{len(list(d.glob('*.png')))} page PNGs in {d}"


def pdf_text(pdf: Path):
    if not PDFTOTEXT:
        return "[pdftotext unavailable]"
    return run([PDFTOTEXT, "-layout", str(pdf), "-"]).stdout


def office_to_pdf(src: Path, out: Path):
    if not Path(SOFFICE).exists():
        return None
    tmp = out / "_pdf"
    tmp.mkdir(parents=True, exist_ok=True)
    run([SOFFICE, "--headless", "--convert-to", "pdf", "--outdir", str(tmp), str(src)])
    p = tmp / f"{src.stem}.pdf"
    return p if p.exists() else None


def docx_text(p: Path):
    import docx
    d = docx.Document(str(p))
    parts = [para.text for para in d.paragraphs]
    for ti, t in enumerate(d.tables, 1):
        parts.append(f"\n[table {ti}]")
        parts += ["\t".join(c.text for c in row.cells) for row in t.rows]
    return "\n".join(parts)


def pptx_text(p: Path):
    from pptx import Presentation
    out = []
    for i, s in enumerate(Presentation(str(p)).slides, 1):
        out.append(f"\n=== slide {i} ===")
        for sh in s.shapes:
            if sh.has_text_frame:
                out.append(sh.text_frame.text)
            if getattr(sh, "has_table", False) and sh.has_table:
                out += ["\t".join(c.text for c in r.cells) for r in sh.table.rows]
        if s.has_notes_slide:
            out.append("[notes] " + s.notes_slide.notes_text_frame.text)
    return "\n".join(out)


def xlsx_text(p: Path, max_rows=2000, max_cols=60):
    import openpyxl
    out = []
    for formulas in (False, True):
        wb = openpyxl.load_workbook(str(p), data_only=not formulas, read_only=True)
        label = "FORMULAS" if formulas else "VALUES"
        for ws in wb.worksheets:
            out.append(f"\n=== [{label}] sheet: {ws.title} ===")
            for r_i, row in enumerate(ws.iter_rows(max_col=max_cols, values_only=True)):
                if r_i >= max_rows:
                    out.append(f"[truncated at {max_rows} rows]")
                    break
                vals = ["" if v is None else str(v) for v in row]
                if any(vals) and (not formulas or any(v.startswith("=") for v in vals)):
                    out.append(f"R{r_i+1}\t" + "\t".join(vals).rstrip())
    return "\n".join(out)


def html_text(p: Path):
    from html.parser import HTMLParser

    class H(HTMLParser):
        def __init__(self):
            super().__init__(); self.buf = []; self.skip = False
        def handle_starttag(self, t, a):
            self.skip = t in ("script", "style")
            if t in ("p", "br", "tr", "li", "h1", "h2", "h3", "div"): self.buf.append("\n")
            if t in ("td", "th"): self.buf.append("\t")
        def handle_endtag(self, t):
            self.skip = False
        def handle_data(self, d):
            if not self.skip: self.buf.append(d)
    h = H(); h.feed(p.read_text(errors="replace"))
    return "".join(h.buf)


def extract(src: Path, out: Path):
    ext = src.suffix.lower()
    note = ""
    if ext == ".pdf":
        text, note = pdf_text(src), render_pdf(src, out)
    elif ext == ".docx":
        text = docx_text(src)
        pdf = office_to_pdf(src, out); note = render_pdf(pdf, out, src.name) if pdf else "no soffice render"
    elif ext == ".pptx":
        text = pptx_text(src)
        pdf = office_to_pdf(src, out); note = render_pdf(pdf, out, src.name) if pdf else "no soffice render"
    elif ext in (".xlsx", ".xlsm"):
        text = xlsx_text(src)
    elif ext in (".csv", ".tsv"):
        with open(src, newline="", errors="replace") as f:
            rows = list(csv.reader(f, delimiter="\t" if ext == ".tsv" else ","))
        text = f"[{len(rows)} rows]\n" + "\n".join("\t".join(r) for r in rows[:5000])
    elif ext in (".html", ".htm"):
        text = html_text(src); note = "HTML: render to PDF/PNG for a visual check"
    elif ext in IMG_EXT:
        text = f"[image] {src} — inspect visually with the Read tool"
    elif ext in TEXT_EXT or ext == "":
        text = src.read_text(errors="replace")
    else:
        text = f"[unsupported format {ext}] — QA metadata/visuals only"
    dst = out / f"{src.name}.txt"
    dst.write_text(text)
    print(f"{src.name}: {len(text):,} chars -> {dst}" + (f" | {note}" if note else ""))


def main():
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    src, out = Path(sys.argv[1]).expanduser(), Path(sys.argv[2]).expanduser()
    out.mkdir(parents=True, exist_ok=True)
    files = [p for p in sorted(src.rglob("*")) if p.is_file()] if src.is_dir() else [src]
    for f in files:
        try:
            extract(f, out)
        except Exception as e:  # keep going; report per-file failure
            print(f"{f.name}: EXTRACT FAILED ({type(e).__name__}: {e})")


if __name__ == "__main__":
    main()
